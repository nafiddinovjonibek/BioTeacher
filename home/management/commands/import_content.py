"""
Kontentni CSV/XLSX fayldan import qilish (FR-68).

Metodist Excel'da kontent tayyorlaydi, dasturchisiz yuklaydi:

    python manage.py import_content --type questions --file savollar.xlsx
    python manage.py import_content --type assignments --file topshiriqlar.csv
    python manage.py import_content --type lessons --file darslar.xlsx --dry-run
    python manage.py import_content --template questions --out namuna.xlsx

Namuna faylni `--template` bilan yarating — ustunlar to'g'ri nomlangan holda chiqadi.
"""

import csv
import io
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

from assignments.models import Assignment, Rubric
from content.models import Lesson, Section, Topic
from core.enums import BloomLevel, Component, Cut, Module
from diagnostics.models import Choice, Question, Questionnaire

# --------------------------------------------------------------------- sxema

SCHEMAS = {
    "questions": {
        "columns": ["questionnaire_slug", "text", "component", "bloom", "reverse",
                    "variant_1", "variant_2", "variant_3", "variant_4", "correct", "explanation"],
        "sample": [
            "test-initial",
            "Fotosintezning yorug'lik bosqichi qayerda kechadi?",
            "COG", "1", "",
            "Tilakoid membranalarida", "Stromada", "Mitoxondriyada", "Sitoplazmada",
            "1",
            "Yorug'lik bosqichi tilakoid membranalarida kechadi.",
        ],
        "help": ("component: MOT/COG/ACT/REF/CRE · bloom: 1-6 · reverse: 1 yoki bo'sh · "
                 "correct: to'g'ri variant raqami (1-4). Likert anketa uchun variantlar "
                 "va correct bo'sh qoldiriladi."),
    },
    "assignments": {
        "columns": ["slug", "module", "kind", "title", "component", "bloom", "difficulty",
                    "minutes", "rubric_slug", "context", "body", "reference"],
        "sample": [
            "lab-yangi-topshiriq", "LAB", "LAB4", "Yangi laboratoriya topshirig'i",
            "COG", "4", "3", "30", "lab-tadqiqot",
            "Vaziyat tavsifi…", "Topshiriq matni…", "Etalon yechim…",
        ],
        "help": ("module: BIOBILIM/LAB/PEDAGOG/RAQAMLI/KREATIV · "
                 "kind: LAB4/LESSON_PLAN/CASE/DIGITAL/CREATIVE · "
                 "rubric_slug mavjud rubrikaga ishora qilishi kerak."),
    },
    "lessons": {
        "columns": ["section_title", "section_slug", "topic_title", "topic_slug",
                    "lesson_title", "lesson_slug", "duration", "body"],
        "sample": [
            "Hujayra biologiyasi", "hujayra-biologiyasi",
            "Hujayra tuzilishi", "hujayra-tuzilishi",
            "Yangi dars", "yangi-dars", "20", "Dars matni…",
        ],
        "help": "Fan va mavzu mavjud bo'lmasa avtomatik yaratiladi.",
    },
}


class Command(BaseCommand):
    help = "Kontentni CSV yoki XLSX fayldan import qiladi (FR-68)."

    def add_arguments(self, parser):
        parser.add_argument("--type", choices=sorted(SCHEMAS), help="Import turi.")
        parser.add_argument("--file", help="CSV yoki XLSX fayl yo'li.")
        parser.add_argument("--template", choices=sorted(SCHEMAS),
                            help="Namuna fayl yaratadi va chiqadi.")
        parser.add_argument("--out", default="namuna.xlsx", help="--template uchun chiqish fayli.")
        parser.add_argument("--dry-run", action="store_true",
                            help="Bazaga yozmasdan tekshiradi.")

    def handle(self, *args, **options):
        if options["template"]:
            return self._write_template(options["template"], options["out"])

        if not options["type"] or not options["file"]:
            raise CommandError("--type va --file majburiy (yoki --template ishlating).")

        path = Path(options["file"])
        if not path.exists():
            raise CommandError(f"Fayl topilmadi: {path}")

        rows = self._read(path, SCHEMAS[options["type"]]["columns"])
        self.stdout.write(f"O'qildi: {len(rows)} qator ({path.name})")

        handler = getattr(self, f"_import_{options['type']}")
        try:
            with transaction.atomic():
                created, updated, errors = handler(rows)
                if options["dry_run"]:
                    transaction.set_rollback(True)
        except Exception as exc:
            raise CommandError(f"Import to'xtatildi: {exc}") from exc

        for error in errors:
            self.stdout.write(self.style.ERROR(f"  {error}"))

        prefix = "[DRY RUN] " if options["dry_run"] else ""
        style = self.style.WARNING if options["dry_run"] else self.style.SUCCESS
        self.stdout.write(style(
            f"\n{prefix}Yaratildi: {created}, yangilandi: {updated}, xato: {len(errors)}"
        ))
        if options["dry_run"]:
            self.stdout.write("Bazaga yozilmadi. Tayyor bo'lsangiz --dry-run ni olib tashlang.")

    # ------------------------------------------------------------- o'qish

    def _read(self, path, expected):
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            rows = self._read_xlsx(path)
        else:
            rows = self._read_csv(path)

        if not rows:
            raise CommandError("Fayl bo'sh yoki sarlavha qatori yo'q.")

        missing = [column for column in expected if column not in rows[0]]
        if missing:
            raise CommandError(
                f"Ustunlar yetishmayapti: {', '.join(missing)}.\n"
                f"Kutilgan tartib: {', '.join(expected)}"
            )
        return rows

    def _read_csv(self, path):
        text = path.read_text(encoding="utf-8-sig")
        # Nuqta-vergul ham, vergul ham qo'llab-quvvatlanadi.
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,") if text.strip() else None
        reader = csv.DictReader(io.StringIO(text), dialect=dialect) if dialect else \
            csv.DictReader(io.StringIO(text), delimiter=";")
        return [{k: (v or "").strip() for k, v in row.items() if k} for row in reader]

    def _read_xlsx(self, path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("XLSX o'qish uchun openpyxl kerak: pip install openpyxl") from exc

        sheet = load_workbook(path, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        result = []
        for raw in rows[1:]:
            if all(cell is None or str(cell).strip() == "" for cell in raw):
                continue
            result.append({
                header[index]: ("" if value is None else str(value).strip())
                for index, value in enumerate(raw) if index < len(header) and header[index]
            })
        return result

    # ---------------------------------------------------------- namuna fayl

    def _write_template(self, kind, out):
        schema = SCHEMAS[kind]
        path = Path(out)
        if path.suffix.lower() == ".csv":
            with path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, delimiter=";")
                writer.writerow(schema["columns"])
                writer.writerow(schema["sample"])
        else:
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font
            except ImportError as exc:
                raise CommandError("openpyxl kerak: pip install openpyxl") from exc

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = kind
            sheet.append(schema["columns"])
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.append(schema["sample"])
            for index in range(1, len(schema["columns"]) + 1):
                sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = 24
            sheet.freeze_panes = "A2"
            workbook.save(path)

        self.stdout.write(self.style.SUCCESS(f"Namuna yaratildi: {path}"))
        self.stdout.write(f"Ustunlar: {', '.join(schema['columns'])}")
        self.stdout.write(self.style.WARNING(schema["help"]))

    # ------------------------------------------------------------- importlar

    def _import_questions(self, rows):
        created = updated = 0
        errors = []
        for number, row in enumerate(rows, start=2):
            try:
                questionnaire = Questionnaire.objects.get(slug=row["questionnaire_slug"])
            except Questionnaire.DoesNotExist:
                errors.append(f"{number}-qator: '{row['questionnaire_slug']}' so'rovnomasi topilmadi")
                continue

            component = row["component"].upper()
            if component not in Component.values:
                errors.append(f"{number}-qator: noma'lum komponent '{component}'")
                continue

            try:
                bloom = int(row["bloom"] or 1)
                if bloom not in BloomLevel.values:
                    raise ValueError
            except ValueError:
                errors.append(f"{number}-qator: bloom 1..6 oralig'ida bo'lishi kerak")
                continue

            question, is_new = Question.objects.update_or_create(
                questionnaire=questionnaire,
                text=row["text"],
                defaults={
                    "component": component,
                    "bloom_level": bloom,
                    "reverse_scored": bool(row.get("reverse", "").strip()),
                    "explanation": row.get("explanation", ""),
                },
            )
            created += is_new
            updated += not is_new

            variants = [row.get(f"variant_{i}", "").strip() for i in range(1, 5)]
            variants = [v for v in variants if v]
            if not variants:
                continue  # Likert savoli

            correct_raw = row.get("correct", "").strip()
            try:
                correct = int(float(correct_raw))
            except ValueError:
                errors.append(f"{number}-qator: 'correct' variant raqami bo'lishi kerak")
                continue
            if not 1 <= correct <= len(variants):
                errors.append(f"{number}-qator: 'correct' 1..{len(variants)} oralig'ida bo'lsin")
                continue

            question.choices.all().delete()
            for index, text in enumerate(variants, start=1):
                Choice.objects.create(
                    question=question, text=text, is_correct=(index == correct), order=index
                )
        return created, updated, errors

    def _import_assignments(self, rows):
        created = updated = 0
        errors = []
        for number, row in enumerate(rows, start=2):
            module = row["module"].upper()
            if module not in Module.values:
                errors.append(f"{number}-qator: noma'lum modul '{module}'")
                continue
            kind = row["kind"].upper()
            if kind not in Assignment.Kind.values:
                errors.append(f"{number}-qator: noma'lum tur '{kind}'")
                continue
            component = row["component"].upper()
            if component not in Component.values:
                errors.append(f"{number}-qator: noma'lum komponent '{component}'")
                continue
            rubric = Rubric.objects.filter(slug=row["rubric_slug"]).first()
            if rubric is None:
                errors.append(f"{number}-qator: '{row['rubric_slug']}' rubrikasi topilmadi")
                continue

            _, is_new = Assignment.objects.update_or_create(
                slug=row["slug"] or slugify(row["title"]),
                defaults={
                    "module": module,
                    "kind": kind,
                    "title": row["title"],
                    "component": component,
                    "bloom_level": int(float(row.get("bloom") or 3)),
                    "difficulty": int(float(row.get("difficulty") or 3)),
                    "estimated_minutes": int(float(row.get("minutes") or 30)),
                    "rubric": rubric,
                    "context_note": row.get("context", ""),
                    "body": row["body"],
                    "reference_solution": row.get("reference", ""),
                    "is_active": True,
                },
            )
            created += is_new
            updated += not is_new
        return created, updated, errors

    def _import_lessons(self, rows):
        created = updated = 0
        errors = []
        for number, row in enumerate(rows, start=2):
            if not row.get("lesson_title"):
                errors.append(f"{number}-qator: lesson_title bo'sh")
                continue

            section, _ = Section.objects.get_or_create(
                slug=row["section_slug"] or slugify(row["section_title"]),
                defaults={"title": row["section_title"]},
            )
            topic, _ = Topic.objects.get_or_create(
                section=section,
                slug=row["topic_slug"] or slugify(row["topic_title"]),
                defaults={"title": row["topic_title"]},
            )
            _, is_new = Lesson.objects.update_or_create(
                topic=topic,
                slug=row["lesson_slug"] or slugify(row["lesson_title"]),
                defaults={
                    "title": row["lesson_title"],
                    "body": row.get("body", ""),
                    "duration_minutes": int(float(row.get("duration") or 20)),
                },
            )
            created += is_new
            updated += not is_new
        return created, updated, errors
